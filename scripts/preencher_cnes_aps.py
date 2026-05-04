from pathlib import Path
import re
import shutil
import unicodedata
import zipfile

import openpyxl
import pandas as pd


BASE_DIR = Path(__file__).resolve().parents[1]
CNES_ZIP = BASE_DIR / "data" / "processado" / "BASE_DE_DADOS_CNES_202603.ZIP"
SAIDA_CSV = BASE_DIR / "data" / "entrada" / "aps_hospitais_cnes.csv"


def achar_planilha():
    for caminho in (BASE_DIR / "referencias").rglob("*.xlsx"):
        if "caps" in caminho.name.lower() and not caminho.name.startswith("~$"):
            return caminho
    raise FileNotFoundError("Planilha endereco_caps.xlsx nao encontrada.")


def normalizar(txt):
    txt = "" if pd.isna(txt) else str(txt)
    txt = unicodedata.normalize("NFKD", txt).encode("ascii", "ignore").decode("ascii")
    txt = re.sub(r"[^A-Za-z0-9]+", " ", txt).upper()
    return re.sub(r"\s+", " ", txt).strip()


def chave_busca(nome):
    termos = normalizar(nome).split()
    return [t for t in termos if t not in {"HOSPITAL", "MUNICIPAL", "SMS", "RIO"}]


def carregar_cnes_rj():
    usecols = [
        "CO_UNIDADE",
        "CO_CNES",
        "NO_FANTASIA",
        "NO_RAZAO_SOCIAL",
        "NO_LOGRADOURO",
        "NU_ENDERECO",
        "NO_BAIRRO",
        "CO_CEP",
        "CO_MUNICIPIO_GESTOR",
        "NU_LATITUDE",
        "NU_LONGITUDE",
        "TP_GESTAO",
        "CO_TIPO_ESTABELECIMENTO",
        "CO_MOTIVO_DESAB",
    ]
    partes = []
    with zipfile.ZipFile(CNES_ZIP) as zf:
        with zf.open("tbEstabelecimento202603.csv") as f:
            for chunk in pd.read_csv(
                f,
                sep=";",
                encoding="latin1",
                dtype=str,
                usecols=lambda c: c in usecols,
                chunksize=200_000,
            ):
                rj = chunk[chunk["CO_UNIDADE"].astype(str).str.startswith("33", na=False)].copy()
                if not rj.empty:
                    partes.append(rj)
        municipios = pd.read_csv(
            zf.open("tbMunicipio202603.csv"),
            sep=";",
            encoding="latin1",
            dtype=str,
            usecols=["CO_MUNICIPIO", "NO_MUNICIPIO"],
        )

    cnes = pd.concat(partes, ignore_index=True).merge(
        municipios,
        left_on="CO_MUNICIPIO_GESTOR",
        right_on="CO_MUNICIPIO",
        how="left",
    )
    cnes["SEARCH"] = (cnes["NO_FANTASIA"].fillna("") + " " + cnes["NO_RAZAO_SOCIAL"].fillna("")).map(normalizar)
    return cnes


def montar_endereco(row):
    partes = [row.get("NO_LOGRADOURO"), row.get("NU_ENDERECO"), row.get("NO_BAIRRO")]
    endereco = ", ".join([str(p).strip() for p in partes if pd.notna(p) and str(p).strip()])
    cep = row.get("CO_CEP")
    if pd.notna(cep) and str(cep).strip():
        endereco = f"{endereco} - CEP {str(cep).strip()}"
    return endereco


def encontrar_hospital(cnes, nome):
    termos = chave_busca(nome)
    if not termos:
        return None
    mask = cnes["CO_MOTIVO_DESAB"].isna() & cnes["TP_GESTAO"].eq("M") & cnes["CO_TIPO_ESTABELECIMENTO"].eq("006")
    for termo in termos:
        mask = mask & cnes["SEARCH"].str.contains(termo, na=False)
    candidatos = cnes[mask].copy()
    if candidatos.empty:
        return None
    return candidatos.iloc[0]


def garantir_colunas(ws, nomes):
    headers = [cell.value for cell in ws[1]]
    for nome in nomes:
        if nome not in headers:
            ws.cell(row=1, column=len(headers) + 1).value = nome
            headers.append(nome)
    return {nome: headers.index(nome) + 1 for nome in headers}


def main():
    planilha = achar_planilha()
    backup = planilha.with_suffix(".backup_antes_cnes.xlsx")
    if not backup.exists():
        shutil.copy2(planilha, backup)

    cnes = carregar_cnes_rj()
    wb = openpyxl.load_workbook(planilha)
    ws = wb["APS"]
    colunas = garantir_colunas(
        ws,
        [
            "CNES",
            "nome_servico_cnes",
            "endereco_cnes",
            "municipio_cnes",
            "lat_cnes",
            "lon_cnes",
            "fonte_cnes",
        ],
    )

    preenchidos = []
    for row_idx in range(2, ws.max_row + 1):
        nome_servico = ws.cell(row=row_idx, column=colunas["nome_servico"]).value
        if not nome_servico:
            continue
        achado = encontrar_hospital(cnes, nome_servico)
        if achado is None:
            continue
        ws.cell(row=row_idx, column=colunas["CNES"]).value = str(achado["CO_CNES"]).zfill(7)
        ws.cell(row=row_idx, column=colunas["nome_servico_cnes"]).value = achado["NO_FANTASIA"]
        ws.cell(row=row_idx, column=colunas["endereco_cnes"]).value = montar_endereco(achado)
        ws.cell(row=row_idx, column=colunas["municipio_cnes"]).value = achado["NO_MUNICIPIO"]
        ws.cell(row=row_idx, column=colunas["lat_cnes"]).value = achado["NU_LATITUDE"]
        ws.cell(row=row_idx, column=colunas["lon_cnes"]).value = achado["NU_LONGITUDE"]
        ws.cell(row=row_idx, column=colunas["fonte_cnes"]).value = "CNES/DataSUS 202603"
        preenchidos.append(
            {
                "nome_servico_planilha": nome_servico,
                "CNES": str(achado["CO_CNES"]).zfill(7),
                "nome_servico_cnes": achado["NO_FANTASIA"],
                "municipio_cnes": achado["NO_MUNICIPIO"],
                "endereco_cnes": montar_endereco(achado),
            }
        )

    wb.save(planilha)
    pd.DataFrame(preenchidos).to_csv(SAIDA_CSV, index=False, encoding="utf-8-sig")
    print(f"Planilha atualizada: {planilha}")
    print(f"Backup: {backup}")
    print(f"Registros APS preenchidos: {len(preenchidos)}")
    print(f"CSV gerado: {SAIDA_CSV}")
    for item in preenchidos:
        print(f"{item['nome_servico_planilha']} -> {item['CNES']} / {item['nome_servico_cnes']}")


if __name__ == "__main__":
    main()
